"""Cache Word page fields, then verify preserved content and final PDF evidence."""
from pathlib import Path
from zipfile import ZipFile, ZIP_DEFLATED
import hashlib
import json
import re
from lxml import etree as E
from docx import Document
from pypdf import PdfReader
import pdfplumber

ROOT=Path(__file__).resolve().parents[1]
BASE=ROOT/'results/study/preserved_report'
DOC=ROOT/'output/docx/thesis_xai_ids_preserved.docx'
W='http://schemas.openxmlformats.org/wordprocessingml/2006/main'

def main():
    fields=json.loads((BASE/'field-values.json').read_text(encoding='utf-8-sig'))
    with ZipFile(DOC) as z:parts={n:z.read(n) for n in z.namelist()}
    tree=E.fromstring(parts['word/document.xml']);updated=0
    for f in tree.iter('{'+W+'}fldSimple'):
        m=re.search(r'PAGEREF (ThesisAnchor\d+)',f.get('{'+W+'}instr',''))
        if m:
            assert m[1] in fields
            assert re.fullmatch(r'[0-9ivxlcdm]+',fields[m[1]],re.I),fields[m[1]]
            list(f.iter('{'+W+'}t'))[0].text=fields[m[1]];updated+=1
    parts['word/document.xml']=E.tostring(tree,encoding='UTF-8',xml_declaration=True,standalone=True)
    with ZipFile(DOC,'w',ZIP_DEFLATED) as z:
        for n,b in parts.items():z.writestr(n,b)
    original=ROOT/'thesis/reference/original_thesis.docx'
    with ZipFile(original) as z:
        excluded={'word/document.xml','word/_rels/document.xml.rels','word/settings.xml'}
        unchanged=[n for n in z.namelist() if n not in excluded]
        assert all(z.read(n)==parts[n] for n in unchanged)
    old=Document(original);new=Document(DOC)
    skip=set(range(52,142))|set(range(143,162))|set(range(163,182))
    text=[p.text for p in new.paragraphs]
    assert all(p.text in text for i,p in enumerate(old.paragraphs) if i not in skip)
    tables=lambda d:[[c.text for r in t.rows for c in r.cells] for t in d.tables]
    assert all(t in tables(new) for t in tables(old))
    captions=[p.text for p in new.paragraphs if p.style.name=='15a Caption-Center' and p.text.startswith('Figure ')]
    assert len(captions)==41
    pdf=ROOT/'output/pdf/thesis_xai_ids_preserved.pdf';reader=PdfReader(pdf)
    alltext='\n'.join(p.extract_text() for p in reader.pages)
    assert 'Error! Reference source not found' not in alltext
    assert 'Error! Bookmark not defined' not in alltext
    for n in range(12,35):assert f'Figure 4.{n}:' in alltext
    outside=[]
    image_count=0
    with pdfplumber.open(pdf) as layout:
        for number,page in enumerate(layout.pages,1):
            for picture in page.images:
                image_count+=1
                if picture['x0'] < -0.5 or picture['top'] < -0.5 or picture['x1'] > page.width+0.5 or picture['bottom'] > page.height+0.5:
                    outside.append(number)
    assert not outside, f'Images outside page bounds: {outside}'
    catalog=json.loads((BASE/'figure_catalog.json').read_text())
    assert len(catalog)==23
    for figure in catalog:
        for extension in ('.png','.svg','.pdf'):
            assert (BASE/'figures'/Path(figure['file']).with_suffix(extension)).is_file()
    bibliography=alltext.split('REFERENCES')[-1].split('APPENDIX A')[0]
    for number in range(1,43):
        assert re.search(r'\['+str(number)+r'\]',bibliography),f'Missing reference {number}'
    report=json.loads((BASE/'preservation_check.json').read_text())
    report.update(pdf_images_checked=image_count,images_outside_page_bounds=outside,separate_figure_formats=['PNG','SVG','PDF'],total_references=42)

    report.update(pdf_pages=len(reader.pages),total_figure_captions=len(captions),total_tables=len(new.tables),cached_page_references=updated,unchanged_package_parts=len(unchanged),status='content, package and page-reference checks passed; visual review recorded separately')
    inputs=[original,DOC,pdf,ROOT/'thesis/manuscript.md',ROOT/'study/preserve_thesis.py',ROOT/'study/preservation_figures.py',ROOT/'study/verify_preserved_thesis.py',ROOT/'results/study/across_seed_summary.csv',ROOT/'results/study/xai/aggregate.csv',ROOT/'results/study/llm/masking_coverage.csv']+list((BASE/'figures').glob('*'))
    report['sha256']={str(p.relative_to(ROOT)).replace('\\','/'):hashlib.sha256(p.read_bytes()).hexdigest() for p in inputs}
    (BASE/'preservation_check.json').write_text(json.dumps(report,indent=2),encoding='utf-8')
    print(json.dumps({k:v for k,v in report.items() if k!='sha256'},indent=2))

if __name__=='__main__':main()
